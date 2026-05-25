"""Build transcript downloads (plain text and Word) without running analysis."""

from __future__ import annotations

import io
import re
from typing import Final

from docx import Document

_SAFE_FILENAME_RE = re.compile(r"[^A-Za-z0-9._\-]+")


def safe_export_stem(filename: str, *, max_len: int = 80) -> str:
    base = (filename or "transcript").rsplit(".", 1)[0].strip() or "transcript"
    s = _SAFE_FILENAME_RE.sub("_", base).strip("._") or "transcript"
    return s[:max_len]


def build_transcript_txt_bytes(full_text: str) -> bytes:
    return (full_text or "").encode("utf-8")


def build_transcript_docx_bytes(*, title: str, full_text: str) -> bytes:
    doc = Document()
    h = (title or "Transcript").strip() or "Transcript"
    doc.add_heading(h, level=0)
    body = (full_text or "").strip()
    for para in body.split("\n\n"):
        p = (para or "").strip()
        if p:
            doc.add_paragraph(p)
    buf: Final[io.BytesIO] = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
