"""Export latest research report to Excel (analysts) and PowerPoint (PR decks)."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.domain.enums import ReportStatus
from app.models.project import Project
from app.models.research_report import ResearchReport

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _as_list(v: Any) -> list[Any]:
    if isinstance(v, list):
        return v
    if isinstance(v, dict):
        return [v]
    return []


def _is_supporting_fact_row(item: dict[str, Any]) -> bool:
    for key in ("entity_type", "type", "unit_type", "entity_kind"):
        raw = item.get(key)
        if isinstance(raw, str) and raw.strip().lower() == "supporting_fact":
            return True
    return False


def _item_text(item: Any) -> str:
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip() if item else ""
    t = item.get("text")
    if isinstance(t, str) and t.strip():
        return t.strip()
    for k in ("finding", "statement", "content", "description"):
        v = item.get(k)
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


def _insight_type_label(item: dict[str, Any]) -> str:
    if _is_supporting_fact_row(item):
        return ""
    for key in ("entity_type", "type", "unit_type"):
        v = item.get(key)
        if isinstance(v, str) and v.strip() and v.strip().lower() != "supporting_fact":
            return v.strip()
    if item.get("finding") and isinstance(item.get("finding"), str):
        return "finding"
    return "key_finding"


def _insight_priority(item: dict[str, Any], rank: int) -> str:
    p = item.get("priority")
    if isinstance(p, (int, float)):
        return str(int(p))
    if isinstance(p, str) and p.strip():
        return p.strip()
    return f"P{rank}"


def _insight_confidence(item: dict[str, Any]) -> str:
    c = item.get("confidence")
    if isinstance(c, (int, float)):
        return f"{float(c):.2f}"
    if isinstance(c, str) and c.strip():
        return c.strip()
    return ""


def _frequency_for_problem(item: dict[str, Any]) -> str:
    g = item.get("grounded_in")
    if isinstance(g, list) and g:
        return str(len(g))
    f = item.get("frequency") or item.get("mentions")
    if isinstance(f, (int, float)):
        return str(int(f))
    if isinstance(f, str) and f.strip():
        return f.strip()
    return ""


def _quotes_for_problem(item: dict[str, Any]) -> str:
    if not isinstance(item, dict):
        return ""
    parts: list[str] = []
    for key in ("comms_signal", "causal_chain", "reputation_gap"):
        v = item.get(key)
        if isinstance(v, str) and v.strip():
            parts.append(v.strip()[:500])
    g = item.get("grounded_in")
    if isinstance(g, list) and g:
        parts.append("Sources: " + ", ".join(str(x) for x in g[:8]))
    return " | ".join(parts) if parts else ""


def _rec_impact_effort(item: dict[str, Any]) -> tuple[str, str]:
    if not isinstance(item, dict):
        return ("", "")
    imp = item.get("impact")
    impact_s = imp.strip() if isinstance(imp, str) else (str(imp) if imp is not None else "")
    eff = item.get("effort")
    if isinstance(eff, str) and eff.strip():
        return (impact_s, eff.strip())
    # Heuristic: PR recommendations often have five structured fields; map tone as "effort" proxy
    tone = item.get("tone_format_how")
    if isinstance(tone, str) and tone.strip():
        return (impact_s, tone.strip()[:240])
    return (impact_s, "")


def _word_frequency_map(extras: dict[str, Any]) -> dict[str, int]:
    wa = extras.get("word_analysis")
    if not isinstance(wa, dict):
        return {}
    wf = wa.get("word_frequency")
    if not isinstance(wf, dict):
        return {}
    out: dict[str, int] = {}
    for k, v in wf.items():
        if isinstance(v, (int, float)):
            out[str(k)] = int(v)
        elif isinstance(v, str) and v.isdigit():
            out[str(k)] = int(v)
    return dict(sorted(out.items(), key=lambda kv: kv[1], reverse=True))


def _autosize_columns(ws: Any, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) // 1 + 2))
        ws.column_dimensions[letter].width = width


def _write_header_row(ws: Any, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_research_report_excel_bytes(project: Project, report: ResearchReport) -> bytes:
    wb = Workbook()
    # Remove default sheet; we'll add named sheets
    default = wb.active
    wb.remove(default)

    # --- Summary ---
    ws_sum = wb.create_sheet("Summary", 0)
    ws_sum["A1"] = "Field"
    ws_sum["B1"] = "Value"
    ws_sum["A1"].font = _HEADER_FONT
    ws_sum["A1"].fill = _HEADER_FILL
    ws_sum["B1"].font = _HEADER_FONT
    ws_sum["B1"].fill = _HEADER_FILL
    rows = [
        ("Project ID", str(project.id)),
        ("Project name", project.name or ""),
        ("Project description", (project.description or "").strip()),
        ("Report title", report.title or ""),
        ("Report status", str(report.status.value if hasattr(report.status, "value") else report.status)),
        ("Updated (UTC)", report.updated_at.astimezone(timezone.utc).isoformat() if report.updated_at else ""),
        ("Executive summary", report.executive_summary or ""),
    ]
    for i, (k, v) in enumerate(rows, start=2):
        ws_sum.cell(row=i, column=1, value=k).alignment = _WRAP
        ws_sum.cell(row=i, column=2, value=v).alignment = _WRAP
    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 90

    extras = report.report_extras_json if isinstance(report.report_extras_json, dict) else {}

    # --- Insights (key findings; no supporting_fact) ---
    ws_ins = wb.create_sheet("Insights")
    ins_headers = ["text", "type", "priority", "confidence"]
    _write_header_row(ws_ins, ins_headers)
    r = 2
    rank = 0
    for it in _as_list(report.key_findings_json):
        if not isinstance(it, dict):
            text = _item_text(it)
            if text:
                ws_ins.cell(row=r, column=1, value=text).alignment = _WRAP
                ws_ins.cell(row=r, column=2, value="key_finding")
                ws_ins.cell(row=r, column=3, value=f"P{r}")
                r += 1
            continue
        if _is_supporting_fact_row(it):
            continue
        rank += 1
        text = _item_text(it)
        if not text:
            continue
        ws_ins.cell(row=r, column=1, value=text).alignment = _WRAP
        ws_ins.cell(row=r, column=2, value=_insight_type_label(it) or "insight").alignment = _WRAP
        ws_ins.cell(row=r, column=3, value=_insight_priority(it, rank)).alignment = _WRAP
        ws_ins.cell(row=r, column=4, value=_insight_confidence(it)).alignment = _WRAP
        r += 1
    _autosize_columns(ws_ins)

    # --- Problems ---
    ws_pr = wb.create_sheet("Problems")
    _write_header_row(ws_pr, ["text", "frequency", "quotes"])
    r = 2
    for it in _as_list(report.problems_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        text = _item_text(it)
        if not text:
            continue
        freq = _frequency_for_problem(it) if isinstance(it, dict) else ""
        quotes = _quotes_for_problem(it) if isinstance(it, dict) else ""
        ws_pr.cell(row=r, column=1, value=text).alignment = _WRAP
        ws_pr.cell(row=r, column=2, value=freq).alignment = _WRAP
        ws_pr.cell(row=r, column=3, value=quotes).alignment = _WRAP
        r += 1
    _autosize_columns(ws_pr)

    # --- Recommendations ---
    ws_rec = wb.create_sheet("Recommendations")
    _write_header_row(ws_rec, ["text", "impact", "effort"])
    r = 2
    for it in _as_list(report.recommendations_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        text = _item_text(it)
        if not text:
            continue
        imp, eff = _rec_impact_effort(it) if isinstance(it, dict) else ("", "")
        ws_rec.cell(row=r, column=1, value=text).alignment = _WRAP
        ws_rec.cell(row=r, column=2, value=imp).alignment = _WRAP
        ws_rec.cell(row=r, column=3, value=eff).alignment = _WRAP
        r += 1
    _autosize_columns(ws_rec)

    # --- Quotes ---
    ws_q = wb.create_sheet("Quotes")
    _write_header_row(ws_q, ["text", "source"])
    r = 2
    for it in _as_list(report.supporting_quotes_json):
        if not isinstance(it, dict):
            continue
        quote = it.get("quote")
        if not isinstance(quote, str) or not quote.strip():
            continue
        note = it.get("note")
        src = it.get("source") or it.get("grounded_in")
        source_parts: list[str] = []
        if isinstance(note, str) and note.strip():
            source_parts.append(note.strip())
        if isinstance(src, list):
            source_parts.append(", ".join(str(x) for x in src[:12]))
        elif isinstance(src, str) and src.strip():
            source_parts.append(src.strip())
        ws_q.cell(row=r, column=1, value=quote.strip()).alignment = _WRAP
        ws_q.cell(row=r, column=2, value=" — ".join(source_parts)).alignment = _WRAP
        r += 1
    _autosize_columns(ws_q)

    # --- Word frequency ---
    ws_wf = wb.create_sheet("Word Frequency")
    _write_header_row(ws_wf, ["word", "count"])
    r = 2
    for word, count in _word_frequency_map(extras).items():
        ws_wf.cell(row=r, column=1, value=word)
        ws_wf.cell(row=r, column=2, value=count)
        r += 1
    _autosize_columns(ws_wf)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _split_bullets(text: str, max_chars: int = 320) -> list[str]:
    """Split long analytical text into slide-friendly bullets."""
    t = text.strip()
    if not t:
        return []
    if "\n\n" in t:
        parts = [p.strip() for p in t.split("\n\n") if p.strip()]
        out: list[str] = []
        for p in parts:
            out.extend(_split_bullets(p.replace("\n", " "), max_chars))
        return out[:24]
    if len(t) <= max_chars:
        return [t]
    sentences = re.split(r"(?<=[.!?])\s+", t)
    chunks: list[str] = []
    cur = ""
    for s in sentences:
        s = s.strip()
        if not s:
            continue
        if len(cur) + len(s) + 1 <= max_chars:
            cur = f"{cur} {s}".strip()
        else:
            if cur:
                chunks.append(cur)
            cur = s
    if cur:
        chunks.append(cur)
    return chunks if chunks else [t[:max_chars] + "…"]


def _add_title_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    *,
    max_bullets: int = 6,
) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    lines = bullets[:max_bullets]
    if not lines:
        p0 = body.paragraphs[0]
        p0.text = "—"
        p0.level = 0
        p0.font.size = Pt(20)
        p0.alignment = PP_ALIGN.LEFT
        return
    for i, line in enumerate(lines):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line[:500]
        p.level = 0
        p.font.size = Pt(20)
        p.alignment = PP_ALIGN.LEFT


def _add_bullet_slides(prs: Presentation, section_title: str, items: list[str], *, per_slide: int = 5) -> None:
    chunk: list[str] = []
    part = 0
    for raw in items:
        for b in _split_bullets(raw, max_chars=280):
            chunk.append(b)
            if len(chunk) >= per_slide:
                part += 1
                suf = f" ({part})" if part > 1 else ""
                _add_title_content_slide(prs, f"{section_title}{suf}", chunk, max_bullets=per_slide)
                chunk = []
    if chunk:
        part += 1
        suf = f" ({part})" if part > 1 else ""
        _add_title_content_slide(prs, f"{section_title}{suf}", chunk, max_bullets=per_slide)


def build_research_report_pptx_bytes(project: Project, report: ResearchReport) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    slide0 = prs.slides.add_slide(prs.slide_layouts[0])
    slide0.shapes.title.text = project.name or "Research report"
    sub_lines = [report.title or ""]
    if report.updated_at:
        sub_lines.append(
            report.updated_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        )
    if len(slide0.placeholders) > 1:
        slide0.placeholders[1].text = "\n".join(s for s in sub_lines if s)

    # 2 Executive summary
    es = (report.executive_summary or "").strip()
    if es:
        bullets = _split_bullets(es, max_chars=400)
        _add_title_content_slide(prs, "Executive summary", bullets[:8], max_bullets=8)

    # 3 Key insights
    insight_lines: list[str] = []
    for it in _as_list(report.key_findings_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        t = _item_text(it)
        if t:
            insight_lines.append(t)
    if insight_lines:
        _add_bullet_slides(prs, "Key insights", insight_lines, per_slide=4)

    # 4 Problems
    prob_lines = []
    for it in _as_list(report.problems_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        t = _item_text(it)
        if t:
            prob_lines.append(t)
    if prob_lines:
        _add_bullet_slides(prs, "Problems", prob_lines, per_slide=5)

    # 5 Patterns
    pat_lines = []
    for it in _as_list(report.patterns_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        t = _item_text(it)
        if t:
            pat_lines.append(t)
    if pat_lines:
        _add_bullet_slides(prs, "Patterns", pat_lines, per_slide=5)

    # 6 Recommendations
    rec_lines = []
    for it in _as_list(report.recommendations_json):
        if isinstance(it, dict) and _is_supporting_fact_row(it):
            continue
        t = _item_text(it)
        if t:
            rec_lines.append(t)
    if rec_lines:
        _add_bullet_slides(prs, "Recommendations", rec_lines, per_slide=4)

    # 7 Quotes
    quote_lines: list[str] = []
    for it in _as_list(report.supporting_quotes_json):
        if not isinstance(it, dict):
            continue
        q = it.get("quote")
        if not isinstance(q, str) or not q.strip():
            continue
        note = it.get("note")
        if isinstance(note, str) and note.strip():
            quote_lines.append(f"«{q.strip()}» — {note.strip()[:120]}")
        else:
            quote_lines.append(f"«{q.strip()}»")
    if quote_lines:
        _add_bullet_slides(prs, "Quotes", quote_lines, per_slide=4)

    # 8 Next steps
    step_lines: list[str] = []
    for it in _as_list(report.next_steps_json):
        if isinstance(it, dict):
            t = str(it.get("text") or "").strip()
            if t:
                step_lines.append(t)
        elif isinstance(it, str) and it.strip():
            step_lines.append(it.strip())
    extras = report.report_extras_json if isinstance(report.report_extras_json, dict) else {}
    for it in _as_list(extras.get("next_steps_pr")):
        if isinstance(it, str) and it.strip():
            step_lines.append(it.strip())
    if step_lines:
        _add_bullet_slides(prs, "Next steps", step_lines, per_slide=6)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def assert_report_exportable(report: ResearchReport | None) -> None:
    if report is None:
        raise ValueError("report_not_found")
    if report.status != ReportStatus.READY:
        raise ValueError("report_not_ready")


def export_filename_stem(project_id: str) -> str:
    return f"research_report_{project_id}_{datetime.now(timezone.utc).strftime('%Y%m%d')}"
