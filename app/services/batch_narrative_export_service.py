"""Narrative PDF/DOCX export from cached full-report payload (not raw JSON)."""

from __future__ import annotations

import os
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.schemas.research_batch import FullReportResponse
from app.services.report_narrative_service import narrative_sections_from_full_report


def resolve_unicode_sans_ttf() -> str | None:
    candidates = [
        os.environ.get("DEJAVU_SANS_TTF"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for p in candidates:
        if p and os.path.isfile(p):
            return p
    return None


def build_narrative_pdf_bytes(full: FullReportResponse, *, batch_name: str, lang: str = "ru") -> bytes:
    sections = narrative_sections_from_full_report(full, lang=lang)
    font_path = resolve_unicode_sans_ttf()
    font_name = "UnicodeSans"
    if font_path and font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, font_path))
    if not font_path:
        font_name = "Helvetica"

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "title",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=16,
        leading=20,
        spaceAfter=6 * mm,
    )
    h_style = ParagraphStyle(
        "h2",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=16,
        spaceAfter=3 * mm,
        spaceBefore=2 * mm,
    )
    body_style = ParagraphStyle(
        "body",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=14,
    )

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=batch_name[:120] or "report",
    )
    story: list = []
    story.append(Paragraph(escape(batch_name or "Report"), title_style))
    story.append(Spacer(1, 4 * mm))
    for sec in sections:
        story.append(Paragraph(escape(sec.title), h_style))
        for block in (sec.body or "").split("\n\n"):
            b = (block or "").strip()
            if not b:
                continue
            safe = escape(b).replace("\n", "<br/>")
            story.append(Paragraph(safe, body_style))
        story.append(Spacer(1, 3 * mm))

    # Evidence appendix (action cards + sample quotes) — decision-ready pack.
    story.append(Paragraph(escape("Evidence appendix"), h_style))
    for q in (full.sample_quotes or [])[:24]:
        line = escape((q.text or "")[:900]).replace("\n", "<br/>")
        if line.strip():
            story.append(Paragraph(line, body_style))
    for card in (full.pr_strategy.action_cards or [])[:20]:
        if str(card.evidence_status) != "ok":
            continue
        hdr = escape((card.action or card.message or "")[:500])
        story.append(Paragraph(f"• {hdr}", body_style))
        for ref in (card.evidence_refs or [])[:5]:
            rq = escape((ref.quote or "")[:500]).replace("\n", "<br/>")
            if rq.strip():
                story.append(Paragraph(f" — {rq}", body_style))
        story.append(Spacer(1, 2 * mm))

    doc.build(story)
    return buf.getvalue()


def build_narrative_docx_bytes(full: FullReportResponse, *, batch_name: str, lang: str = "ru") -> bytes:
    sections = narrative_sections_from_full_report(full, lang=lang)
    document = Document()
    n_style = document.styles["Normal"]
    n_style.font.size = Pt(11)
    document.add_heading(batch_name or "Report", level=0)
    for sec in sections:
        document.add_heading(sec.title, level=1)
        for block in (sec.body or "").split("\n\n"):
            b = (block or "").strip()
            if b:
                document.add_paragraph(b)
    bio = BytesIO()
    document.save(bio)
    return bio.getvalue()
