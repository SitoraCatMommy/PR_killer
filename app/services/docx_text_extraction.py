"""Extract plain text from Word .docx (OOXML) for research ingestion."""

from __future__ import annotations

import logging
from io import BytesIO
from zipfile import BadZipFile

from docx import Document

logger = logging.getLogger(__name__)

_DOCX_OPEN_ERRORS: tuple[type[BaseException], ...] = (BadZipFile, ValueError, KeyError)
try:
    from docx.opc.exceptions import PackageNotFoundError as _DocxPackageNotFoundError

    _DOCX_OPEN_ERRORS = (*_DOCX_OPEN_ERRORS, _DocxPackageNotFoundError)
except ImportError:
    pass

DOCX_ERROR_CODES = frozenset(
    {
        "docx_empty",
        "docx_corrupt_or_unreadable",
        "doc_format_not_supported_use_docx",
    },
)

DOCX_ERROR_MESSAGES_EN: dict[str, str] = {
    "docx_empty": "The Word file contains no extractable text.",
    "docx_corrupt_or_unreadable": "The Word file could not be read. Make sure it is a valid .docx file.",
    "doc_format_not_supported_use_docx": "Legacy .doc format is not supported. Save the document as .docx.",
}

_DOCX_MIME_PREFIXES = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
)


def is_docx_filename_or_mime(filename: str, mime_type: str | None) -> bool:
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    if ext == "docx" or (filename or "").lower().endswith(".docx"):
        return True
    m = (mime_type or "").lower().split(";")[0].strip()
    return m.startswith(_DOCX_MIME_PREFIXES)


def is_legacy_doc_filename_or_mime(filename: str, mime_type: str | None) -> bool:
    if (filename or "").lower().endswith(".doc") and not (filename or "").lower().endswith(".docx"):
        return True
    m = (mime_type or "").lower().split(";")[0].strip()
    return m == "application/msword"


def extract_plain_text_from_docx(content: bytes) -> str:
    """Return UTF-8 text from paragraphs and tables. Raises ``ValueError`` with a stable code."""
    if not content:
        raise ValueError("docx_corrupt_or_unreadable")
    try:
        doc = Document(BytesIO(content))
    except _DOCX_OPEN_ERRORS as e:
        logger.info("docx_open_failed bytes=%s err=%s", len(content), type(e).__name__)
        raise ValueError("docx_corrupt_or_unreadable") from e
    except OSError as e:
        logger.info("docx_open_failed_os bytes=%s err=%s", len(content), e)
        raise ValueError("docx_corrupt_or_unreadable") from e

    chunks: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            chunks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            row_txt = " | ".join(c for c in cells if c)
            if row_txt:
                chunks.append(row_txt)

    text = "\n".join(chunks).strip()
    if not text:
        raise ValueError("docx_empty")
    logger.debug("docx_extract_ok chars=%s", len(text))
    return text
